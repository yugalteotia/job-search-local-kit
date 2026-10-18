"""Materialized spreadsheet and HTML/JSON views from canonical SQLite rows."""
import json
import os
from pathlib import Path
import tempfile

from local_paths import ROOT
import job_db

def regenerate(root=ROOT,database=None):
    root=Path(root);database=Path(database or root/'.workflow/jobs.sqlite3')
    import sys
    sys.path.insert(0,str(root/'.local-deps'))
    from openpyxl import load_workbook,Workbook
    db=job_db.connect(database)
    try:rows=job_db.tracker_rows(db)
    finally:db.close()
    report=root/'.workflow/migration/verification.json'
    columns=json.loads(report.read_text())['columns'] if report.is_file() else list(dict.fromkeys(k for r in rows for k in r))
    backup=root/'.workflow/migration/tracker-before-sqlite.xlsx'
    wb=load_workbook(backup) if backup.is_file() else Workbook()
    if 'Tracker' in wb:
        ws=wb['Tracker'];ws.delete_rows(1,ws.max_row)
    else:
        ws=wb.active;ws.title='Tracker'
    ws.append(columns)
    for row in rows:ws.append([row.get(c) for c in columns])
    ws.auto_filter.ref=ws.dimensions
    ws.freeze_panes='A2'
    target=root/'tracker.xlsx';target.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(dir=target.parent,prefix='.tracker-view-',suffix='.xlsx');os.close(fd)
    try:wb.save(tmp);os.replace(tmp,target)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)
    # HTML and JSON are rendered directly from the SQLite rows, not re-read from Excel.
    from tracker_views import views
    class Context:
        ROOT=root
        STATE=root/'.workflow'
        read_json=staticmethod(lambda p,default=None:json.loads(Path(p).read_text()) if Path(p).is_file() else default)
        write_json=staticmethod(lambda p,value:Path(p).write_text(json.dumps(value,indent=2,ensure_ascii=False,default=str)+'\n'))
        key=staticmethod(lambda row:job_db.identity(row.get('Company'),row.get('Job ID'),row.get('Direct link'),row.get('Role')))
    output=views(Context(),rows_override=rows,columns_override=columns)
    # Cross-check generated workbook against DB at the point of generation.
    check=load_workbook(target,read_only=True,data_only=True)['Tracker']
    if check.max_row-1!=len(rows):raise ValueError('Generated tracker row count differs from SQLite')
    result={'database_rows':len(rows),'tracker_rows':check.max_row-1,**output}
    try:
        import quick_status,job_config
        result['status']=quick_status.write(root,job_config.load(),database)
    except Exception as exc:result['status_error']=type(exc).__name__
    return result
