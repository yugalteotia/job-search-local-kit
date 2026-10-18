"""Local CLI for SQLite-backed discovery-only scheduling and guarded legacy commands."""
from local_paths import ROOT
import argparse, contextlib, datetime as dt, fcntl, hashlib, json, os, re, shutil, subprocess, sys, uuid
from pathlib import Path
from zoneinfo import ZoneInfo
import openpyxl
from urllib.parse import urlsplit

STATE = ROOT / '.workflow'
MASTER = ROOT / 'master-resume/Yugal_Teotia_CV.tex'
PROTECTED = ('applied', 'submission unconfirmed', 'rejected', 'interview', 'offer', 'withdrawn', 'referred', 'stopped')

def now():
    return dt.datetime.now(dt.timezone.utc).isoformat()

def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    tmp.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str) + '\n')
    tmp.replace(path)

def read_json(path, default=None):
    return json.loads(path.read_text()) if path.exists() else default

def event(event_type, **data):
    STATE.mkdir(exist_ok=True)
    with (STATE / 'events.jsonl').open('a') as f:
        f.write(json.dumps({'time': now(), 'event': event_type, **data}, default=str) + '\n')

def check_pause():
    if read_json(STATE / 'control.json', {}).get('paused'):
        raise ValueError('Paused. Run resume before further workflow mutations.')

def normalize(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())

def job_folder(slug):
    if not re.fullmatch(r'[a-z0-9-]+', slug):
        raise ValueError('Unsafe job slug')
    bound=os.environ.get('JOB_SEARCH_WORKER_SLUG')
    if bound and slug!=bound:
        raise ValueError('Worker cannot access another job folder')
    return ROOT / 'jobs' / slug

def worker_url_guard(url):
    original=os.environ.get('JOB_SEARCH_WORKER_URL')
    if not original:return
    candidate=urlsplit(str(url));source=urlsplit(original)
    if candidate.scheme!='https' or not candidate.hostname or not source.hostname:
        raise ValueError('Worker browser evidence must use an HTTPS application URL')
    # Exact original host, or an observed continuation under the same ATS domain.
    if candidate.hostname!=source.hostname and candidate.hostname.split('.')[-2:]!=source.hostname.split('.')[-2:]:
        raise ValueError('Worker application URL crossed to a different ATS; handoff required')

def fresh(timestamp):
    checked = dt.datetime.fromisoformat(timestamp.replace('Z', '+00:00'))
    if checked.tzinfo is None:
        raise ValueError('Capture timestamp must have a timezone')
    age = (dt.datetime.now(dt.timezone.utc) - checked).total_seconds()
    if age < -300 or age > 86400:
        raise ValueError('Verification capture stale or future-dated: reverify original; posting age alone does not exclude the role')

def employer(s):
    n = normalize(s)
    return 'unitedhealthgroup' if 'optum' in n or 'unitedhealthgroup' in n else n

def key(row):
    company = employer(row.get('Company'))
    jobid = normalize(row.get('Job ID'))
    return (company, jobid) if jobid else (company, normalize(row.get('Direct link')) or normalize(row.get('Role')))

@contextlib.contextmanager
def locked():
    STATE.mkdir(exist_ok=True)
    with (STATE / 'lock').open('a') as f:
        fcntl.flock(f, fcntl.LOCK_EX)
        yield

def tracker_rows():
    wb = openpyxl.load_workbook(ROOT / 'tracker.xlsx')
    ws = wb['Tracker']
    cols = [c.value for c in ws[1]]
    return wb, ws, cols, [dict(zip(cols, row)) for row in ws.iter_rows(min_row=2, values_only=True)]

def upsert(data, reason, allow_paused=False, authorized_transition=False):
    if (ROOT/'.workflow/canonical-db-enabled.json').is_file():
        raise ValueError('Legacy tracker upsert disabled: SQLite state machine is canonical')
    if not allow_paused:
        check_pause()
    bound_req=os.environ.get('JOB_SEARCH_WORKER_REQ')
    if bound_req:
        if key(data)!=(os.environ.get('JOB_SEARCH_WORKER_COMPANY'),normalize(bound_req)):
            raise ValueError('Worker tracker update is not bound to its requisition')
        if data.get('Direct link') and data['Direct link']!=os.environ.get('JOB_SEARCH_WORKER_URL'):
            raise ValueError('Worker tracker URL mismatch')
    with locked():
        wb, ws, cols, rows = tracker_rows()
        unknown = set(data) - set(cols)
        if unknown:
            raise ValueError(f'Unknown tracker columns: {unknown}')
        matches = [i for i, r in enumerate(rows, 2) if key(r) == key(data)]
        if len(matches) > 1:
            raise ValueError('Existing duplicate requisitions: resolve manually')
        index = matches[0] if matches else ws.max_row + 1
        before = rows[index - 2] if matches else {}
        old = str(before.get('Status', '')).lower()
        import terminal_cleanup as tc
        if tc.is_terminal(old) and any(k not in tc.MINIMAL for k in data):
            raise ValueError('Terminal history row is minimal and cannot gain bulky fields')
        if any(p in old for p in PROTECTED) and 'Status' in data and data.get('Status') != before.get('Status'):
            auth=read_json(STATE/'optum-2388999-authorization.json',{})
            submission_transition=(old in ('stopped','submission unconfirmed') and data.get('Status') in ('Submission unconfirmed','Applied'))
            response_transition=(old=='applied' and data.get('Status')=='Rejected' and data.get('Confirmation evidence') and (ROOT/data['Confirmation evidence']).is_file())
            scoped=(authorized_transition and auth.get('batch_confirmed') is True and auth.get('submission_authorized_after_confirmed_batch') is True and key(data)==('unitedhealthgroup','2388999') and (submission_transition or response_transition))
            daily=read_json(STATE/'daily-authorization.json',{}).get('authorized') is True
            daily_confirm=(daily and authorized_transition and old=='submission unconfirmed' and data.get('Status')=='Applied' and data.get('Confirmation evidence') and (ROOT/data['Confirmation evidence']).is_file() and key(data)!=('unitedhealthgroup','2388999'))
            daily_response=(daily and authorized_transition and old in ('applied','interview','offer') and data.get('Status') in ('Rejected','Interview','Offer','Withdrawn') and data.get('Confirmation evidence') and (ROOT/data['Confirmation evidence']).is_file())
            scoped=scoped or daily_confirm or daily_response
            if not scoped:
                raise ValueError('Protected application state cannot be overwritten by discovery')
        if data.get('Status') == 'Applied' and not data.get('Confirmation evidence'):
            raise ValueError('Applied requires confirmation evidence')
        changes = {k: {'before': before.get(k), 'after': v} for k, v in data.items() if before.get(k) != v}
        if not changes:
            return index
        stamp = dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S') + '-' + uuid.uuid4().hex[:8]
        backup = STATE / 'backups' / (stamp + '.xlsx')
        backup.parent.mkdir(exist_ok=True)
        shutil.copy2(ROOT / 'tracker.xlsx', backup)
        for c, name in enumerate(cols, 1):
            if name in data:
                value = data[name]
                cell = ws.cell(index, c, value=value)
                if isinstance(value, str):
                    cell.data_type = 's'  # never interpret captured text as an Excel formula
        ws.auto_filter.ref = ws.dimensions
        tmp = STATE / 'tracker.tmp.xlsx'
        wb.save(tmp)
        tmp.replace(ROOT / 'tracker.xlsx')
        write_json(STATE / 'changes' / (stamp + '.json'), {'row': index, 'reason': reason, 'changes': changes, 'backup': str(backup.relative_to(ROOT))})
        event('tracker_updated', row=index, reason=reason, changes=changes)
        return index

def plain_tex(text):
    text = re.sub(r'\\(?:textbf|textit)\{([^{}]*)\}', r'\1', text)
    return text.replace('\\%', '%').replace('\\&', '&').replace('~', ' ')

def master_bullets():
    source = MASTER.read_text()
    blocks = re.findall(r'\\begin\{itemize\}(.*?)\\end\{itemize\}', source, re.S)
    return [[plain_tex(b.strip()) for b in re.split(r'\\item\s+', block)[1:]] for block in blocks]

def audit(as_of=None):
    source = MASTER.read_text()
    profile = (ROOT / 'APPLICATION_PROFILE.md').read_text()
    today = as_of or dt.datetime.now(ZoneInfo('Asia/Kolkata')).date()
    release = dt.date(2026,9,30)
    joining = dt.date(2026,10,1)
    fields = {
        'name': ('Yugal Teotia', 'master header'), 'email': ('teotia.yugal@gmail.com', 'master header'),
        'phone': ('+91 9560369538', 'master header'), 'location': ('Pune, India', 'master header'),
        'linkedin': ('https://www.linkedin.com/in/yugalteotia/', 'master header'),
        'github': ('https://github.com/yugalteotia', 'master header'),
    }
    answers = {k: {'value': v, 'state': 'confirmed' if v.rstrip('/') in source else 'uncertain', 'evidence': e} for k, (v, e) in fields.items()}
    answers['first_name'] = {'value':'Yugal','state':answers['name']['state'],'evidence':'First component of master full name'}
    answers['last_name'] = {'value':'Teotia','state':answers['name']['state'],'evidence':'Last component of master full name'}
    answers['country'] = {'value':'India','state':answers['location']['state'],'evidence':'Current location in master, not citizenship/work authorization'}
    for k, value in [('relocation', 'Any city in India'), ('work_mode', 'Onsite, hybrid or remote')]:
        answers[k] = {'value': value, 'state': 'confirmed', 'evidence': 'Both profile documents; preference, not a resume claim'}
    joining_confirmed = '1 October 2026' in profile and '"its 1 oct i am ready"' in profile
    answers['joining_date'] = {'value':joining.isoformat(),'state':'confirmed' if joining_confirmed else 'uncertain','evidence':'Explicit user confirmation preserved in APPLICATION_PROFILE.md (2026-09-27); stale SEARCH_PROFILE line reconciled'}
    answers['total_experience_approx'] = {'value':'2+ years overall professional experience','state':'confirmed','evidence':'Master employment month ranges: Oct 2024–Sep 2026 plus Feb–May 2023; rounded overall claim only, excludes training'}
    answers['java_professional_experience_approx'] = {'value':'Approximately 2 years at IDeaS','state':'confirmed','evidence':'Master IDeaS Oct 2024–Sep 2026; not a claim of 2+ professional Java years; PG-DAC separate'}
    answers['total_experience'] = {'value':None,'state':'uncertain','evidence':'Exact numeric months/decimal years need exact employment dates; month-only dates support rounded summaries only'}
    answers['java_professional_experience'] = {'value':None,'state':'uncertain','evidence':'Exact IDeaS start date needed for numeric professional Java experience; training excluded'}
    dates = read_json(ROOT / 'employment_dates.json', {})
    if dates.get('employment'):
        totals = {'total_experience': [0,0], 'java_professional_experience':[0,0]}
        incomplete = set()
        for employment in dates['employment']:
            if not employment.get('start') or not employment.get('end'):
                incomplete.add('total_experience')
                if employment['java']:
                    incomplete.add('java_professional_experience')
                answers[normalize(employment['company'])+'_employment_dates'] = {'value':None,'state':'uncertain','evidence':employment['basis']}
                answers[normalize(employment['company'])+'_employment_months'] = {'value':{'start':employment.get('start_month'),'end':employment.get('end_month')},'state':'confirmed','evidence':employment['basis']+' Use only in month/year fields.'}
                continue
            start = dt.date.fromisoformat(employment['start'])
            stop = min(today, dt.date.fromisoformat(employment['end']) + dt.timedelta(days=1))
            if stop < start:
                continue
            months = (stop.year-start.year)*12+stop.month-start.month-(stop.day<start.day)
            for field in ['total_experience'] + (['java_professional_experience'] if employment['java'] else []):
                totals[field][0] += months
                totals[field][1] += (stop-start).days
            answers[normalize(employment['company'])+'_employment_dates'] = {'value':{'start':employment['start'],'end':employment['end']},'state':'confirmed','evidence':employment['basis']+'; portal-only, not resume day dates'}
        for field, (months,days) in totals.items():
            if field in incomplete:
                answers[field] = {'value':None,'state':'uncertain','evidence':'Data Nova exact dates are not established; user disallowed assumed dates and the 27-month figure as a confirmed answer. Ask only when a portal requires exact dates or exact total experience.'}
            else:
                answers[field] = {'value':{'completed_months':months,'years':months//12,'months':months%12,'elapsed_days':days,'decimal_years':round(days/365.2425,3)},'state':'confirmed','evidence':'Derived from exact IDeaS dates in employment_dates.json; full calendar months; decimal years = elapsed days / 365.2425. Past end dates inclusive, current day excluded; training excluded.'}
    reason = ("My role at IDeaS ends on 30 September 2026, and I'm seeking a backend-focused software engineering position where I can build services and APIs." if today<=release else "My role at IDeaS ended, and I'm seeking a backend-focused software engineering position where I can build services and APIs.")
    answers['reason_for_leaving'] = {'value':reason,'state':'confirmed' if reason in profile else 'uncertain','evidence':'Confirmed reason in both profiles; tense follows confirmed last working day; no new reason inferred'}
    answers['notice_period'] = {'value':'Last working day 30 September 2026; available from 1 October 2026' if today<joining else 'No remaining notice; available to join','state':'confirmed' if joining_confirmed else 'uncertain','evidence':'Release/availability dates; use only in an open-text availability or remaining-notice question'}
    answers['remaining_notice_days'] = {'value':max(0,(joining-today).days),'state':'confirmed' if joining_confirmed else 'uncertain','evidence':f'Calendar days from {today.isoformat()} IST to confirmed joining date; only for days until joining'}
    answers['contractual_notice_period'] = {'value':None,'state':'uncertain','evidence':'Original employment-contract notice duration is not in the documents; ask only if a form explicitly requires it'}
    for k in ['citizenship', 'dob', 'address', 'education_grade']:
        answers[k] = {'value': None, 'state': 'uncertain', 'evidence': 'TBD in application profile'}
    answers['gender'] = {'value':'Male' if 'Gender: Male' in profile else None,'state':'confirmed' if 'Gender: Male' in profile else 'uncertain','evidence':'Latest explicit user confirmation in APPLICATION_PROFILE.md'}
    answers['work_authorization'] = {'value':'Unrestricted authorization to work in India' if 'unrestricted authorization to work in India' in profile else None,'state':'confirmed' if 'unrestricted authorization to work in India' in profile else 'uncertain','evidence':'User confirmation, India jobs only; not citizenship'}
    for k,value,phrase in [('sponsorship_india','No','no sponsorship required for an India-based role'),('armed_forces','No','No Indian Armed Forces service'),('restrictive_covenant','No','No restrictive covenant'),('prior_uhg_offer','No','no prior UHG offer'),('prior_uhg_employment','No','no prior UHG/affiliate employment')]:
        answers[k]={'value':value,'state':'confirmed' if phrase in profile else 'uncertain','evidence':'Explicit user confirmation in APPLICATION_PROFILE.md; reuse only for an equivalent question'}
    for k,value in [('current_annual_total_ctc_inr',1000000),('current_annual_fixed_pay_inr',900000),('current_annual_variable_pay_inr',100000),('data_nova_annual_total_ctc_inr',360000)]:
        answers[k]={'value':value,'state':'confirmed','mandatory_only':True,'evidence':'Explicit user salary confirmation; exact label and units must match'}
    for k, value in [('current_ctc', '10 LPA annual total CTC (9 LPA fixed + 1 LPA performance-based variable)'), ('expected_ctc', 'Approximately 15 LPA')]:
        answers[k] = {'value': value, 'state': 'confirmed', 'mandatory_only': True, 'evidence': 'Both profile documents; user-provided salary, not independently verified'}
    report = {'checked': now(), 'as_of_ist':today.isoformat(), 'master_sha256': hashlib.sha256(source.encode()).hexdigest(), 'answers': answers,
              'notes': ['Generic master remains source of truth; employment tense follows confirmed end date.', 'Legacy tailored resume contains unsupported claims; do not reuse.', 'Data Nova Kotlin/Agile claims in legacy helper are not established by master experience bullets.', 'Daily submissions require passed private browser probe, active control gate and pre-submit verification. Optum 2388999 never retried.']}
    write_json(STATE / 'profile-audit.json', report)
    event('profile_audited', uncertain=[k for k,v in answers.items() if v['state'] != 'confirmed'])
    return report

def score(packet):
    # Deliberately evidence-led: no keyword-only High verdict.
    if packet.get('live') is False:
        return 'Excluded', 0
    requirements = packet.get('requirements', [])
    if not requirements:
        return 'Unrated', 0
    for req in requirements:
        if req.get('basis') == 'work':
            index = req.get('master_bullet')
            if not isinstance(index, list) or len(index) != 2:
                raise ValueError('Work evidence needs [experience block, bullet]')
            req['evidence'] = master_bullets()[index[0]][index[1]]
        elif req.get('basis') not in ('project', 'education', 'gap', 'profile') or not req.get('evidence'):
            raise ValueError('Every requirement needs evidence or an explicit gap')
    required = [r for r in requirements if r.get('required')]
    covered = sum(r['basis'] in ('work', 'education', 'profile') for r in required)
    percent = round(100 * covered / max(1, len(required)))
    checks=packet.get('fit_checks',{})
    substantive_fit=all(isinstance(checks.get(k),dict) and checks[k].get('verified') is True and checks[k].get('evidence') for k in ('java_spring_backend','experience_band','india_location'))
    high = (packet.get('original_verified') is True and packet.get('live') is True and packet.get('india') is True
            and packet.get('development_role') is True and substantive_fit and required and covered == len(required))
    return ('High' if high else 'Medium' if percent >= 50 else 'Low'), percent

def ingest(path):
    packet=read_json(path,{})
    candidate={'Company':packet.get('company'),'Job ID':packet.get('job_id'),'Role':packet.get('role'),'Direct link':packet.get('url')}
    lock_name=hashlib.sha256(('\0'.join(key(candidate))).encode()).hexdigest()+'.lock'
    lock_path=STATE/'ingest-claims'/lock_name;lock_path.parent.mkdir(parents=True,exist_ok=True)
    with lock_path.open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX)
        return _ingest_locked(path)

def _ingest_locked(path):
    check_pause()
    p = read_json(path)
    for field in ['company', 'role', 'url', 'checked_at', 'jd', 'slug']:
        if not p.get(field):
            raise ValueError(f'Missing {field}')
    if not re.fullmatch(r'[a-z0-9-]+', p['slug']):
        raise ValueError('Unsafe slug')
    candidate={'Company':p['company'],'Job ID':p.get('job_id'),'Role':p['role'],'Direct link':p['url']}
    import terminal_cleanup as tc
    existing_same=[r for r in tracker_rows()[3] if key(r)==key(candidate)]
    if any(tc.is_terminal(r.get('Status')) for r in existing_same):
        raise ValueError('Terminal requisition already in tracker; never resubmit or recreate bulky job files')
    if any(any(word in str(r.get('Status','')).lower() for word in PROTECTED) for r in existing_same):
        raise ValueError('Existing or uncertain application requisition; do not requeue or overwrite')
    if any(key(r)==key(candidate) for r in read_json(STATE/'cleanup/other_history_keys.json',[])):
        raise ValueError('Historical application requisition already recorded; never reapply')
    fit, percent = score(p)
    if fit == 'High':
        if not p.get('job_id'):
            raise ValueError('Employer requisition ID required before High-fit parallel work')
        import priority_queue
        # Verified essentials dominate; posting recency orders otherwise strong matches.
        try:
            posted=dt.datetime.fromisoformat(str(p.get('posted_date','')).replace('Z','+00:00'))
            if posted.tzinfo is None:posted=posted.replace(tzinfo=dt.timezone.utc)
            age_hours=max(0,(dt.datetime.now(dt.timezone.utc)-posted).total_seconds()/3600)
            recency=max(0,20-min(20,age_hours/12))
        except (ValueError,TypeError):recency=0
        priority_queue.enqueue(p['slug'],p['company'],p['job_id'],p['url'],min(100,percent*.8+recency),p['checked_at'])
    folder = job_folder(p['slug'])
    folder.mkdir(exist_ok=True)
    write_json(folder / 'listing.json', {**p, 'fit': fit, 'score': percent})
    (folder / 'JD.md').write_text(f"# {p['company']} — {p['role']}\n\nSource: {p['url']}\nChecked: {p['checked_at']}\nPosting date: {p.get('posted_date') or 'Unknown'}\n\n{p['jd']}\n")
    gaps = [r['name'] for r in p.get('requirements', []) if r['basis'] in ('gap', 'project')]
    row = {'Company': p['company'], 'Role': p['role'], 'Job ID': p.get('job_id'), 'Direct link': p['url'],
           'Source': p.get('source', 'Playwright capture'), 'Location': p.get('location'), 'Posted date': p.get('posted_date'),
           'Verified date': p['checked_at'] if p.get('original_verified') else None, 'Fit': fit,
           'Fit rationale': p.get('rationale', '') + f' Required evidence coverage: {percent}%.',
           'Key gaps': '; '.join(gaps), 'Status': 'Excluded — closed' if fit=='Excluded' else 'Shortlisted', 'Next step': ('Process High-fit role in this run; verify history, resume, facts and submission proof' if read_json(STATE/'control.json',{}).get('submission_enabled') else 'Prepare; submissions disabled') if fit == 'High' else 'Closed; do not apply' if fit=='Excluded' else 'Resolve gaps / verify original before preparation'}
    # Rediscovery must preserve existing application state.
    _, _, _, existing = tracker_rows()
    matches = [r for r in existing if key(r) == key(row)]
    if matches:
        row['Status'] = matches[0]['Status']
        row['Next step'] = matches[0]['Next step']
    upsert(row, 'listing ingested')
    (folder / 'record.md').write_text(f"# {p['company']} — {p['role']}\n\nRequisition: {p.get('job_id') or 'Unknown'}\nFit: {fit} ({percent}% required evidence coverage)\n{p.get('rationale', '')}\n\nGaps: {'; '.join(gaps)}\n\nSee listing.json for requirement-by-requirement evidence; tracker.xlsx for current state; .workflow/events.jsonl for audit.\n")
    event('listing_ingested', slug=p['slug'], fit=fit)
    return p['slug']

def tailor(slug):
    check_pause()
    folder = job_folder(slug)
    listing = read_json(folder / 'listing.json')
    identity={'Company':listing.get('company'),'Job ID':listing.get('job_id'),'Role':listing.get('role'),'Direct link':listing.get('url')}
    if key(identity)==('unitedhealthgroup','2388999'):
        raise ValueError('Optum 2388999 is completed and Rejected; never modify or retry')
    if any(key(r)==key(identity) for r in read_json(STATE/'cleanup/other_history_keys.json',[])):
        raise ValueError('Historical application requisition already recorded; never reapply')
    existing=next((r for r in tracker_rows()[3] if key(r)==key(identity)),{})
    if any(p in str(existing.get('Status','')).lower() for p in PROTECTED):
        raise ValueError('Do not tailor an already actioned or uncertain application')
    if not listing or listing['fit'] != 'High':
        raise ValueError('Only verified High-fit listings may be tailored')
    fresh(listing['checked_at'])
    b = master_bullets()
    today = dt.date.today()
    spec = {
      'summary': 'Backend-focused Software Engineer with 2+ years of hands-on Java development, including PG-DAC training. Experienced in REST and gRPC APIs, Spring Boot, MongoDB, and Ember.js/TypeScript SaaS features. Uses AI-assisted engineering while retaining ownership of design and validation.',
      'skills': [{'label':'Backend / production','items':'Java 17, Spring Boot, Spring MVC, Spring Security, REST APIs, gRPC, Protocol Buffers, MongoDB'},
                 {'label':'Full stack / engineering','items':'Ember.js, TypeScript, Git, Gradle, Maven, GitHub Actions, Docker, Linux, JUnit, Mockito'},
                 {'label':'Project / training','items':'SQL, MySQL, Spring Data JPA, Hibernate, Flyway, OpenAPI; Kotlin, Jetpack Compose; React.js (basic)'},
                 {'label':'AI-assisted engineering','items':'GitHub Copilot, agent workflows, RCA, debugging, scenario generation'}],
      'experience':[{'title':'Software Engineer','org':'IDeaS - A SAS Company','location':'Pune, India','dates':'Oct 2024 - Present' if today <= dt.date(2026,9,30) else 'Oct 2024 - Sep 2026','bullets':[b[0][i] for i in [0,2,3,4,5]]},
                    {'title':'Application Developer','org':'Data Nova India Pvt. Ltd.','location':'Noida, India','dates':'Feb 2023 - May 2023','bullets':b[1][:2]}],
      'projects':[{'title':'UNITY - Peer-to-Peer Music Sync','note':'Functional prototype; AI-assisted','bullets':b[2]},
                  {'title':'Store REST API','note':'In progress; project skills','bullets':b[3]}],
      'education':[{'title':'PG Diploma in Advanced Computing (PG-DAC)','org':'C-DAC','location':'Pune','dates':'Mar 2024 - Aug 2024'},
                   {'title':'B.Tech, Computer Science & Engineering','org':'Dr. A.P.J. Abdul Kalam Technical University','dates':'2020 - 2023'}]}
    write_json(folder / 'resume_spec.json', spec)
    write_json(folder / 'resume_provenance.json', {'master':str(MASTER.relative_to(ROOT)), 'sha256':hashlib.sha256(MASTER.read_bytes()).hexdigest(), 'experience_bullets':'Verbatim from master; subset reordered', 'summary':'Condensed generic summary', 'skills':'Master skills with project/training distinctions from profiles', 'listing':listing['url']})
    from render_resume import build
    from resume_naming import employer_filename
    out = folder / employer_filename(listing['role'], listing['company'], listing.get('job_id'))
    if out.exists():
        archive = folder / 'resume_versions'
        archive.mkdir(exist_ok=True)
        old_hash = hashlib.sha256(out.read_bytes()).hexdigest()
        shutil.copy2(out, archive / (old_hash + '.pdf'))
    build(spec, str(out))
    # Rendering invalidates any earlier review, even if extracted text is identical.
    write_json(folder/'resume_visual_review.json',{'status':'needs_visual_review','pdf':out.name})
    write_json(folder/'resume_latest.json',{'pdf':out.name,'pdf_sha256':hashlib.sha256(out.read_bytes()).hexdigest()})
    subprocess.run(['pdftoppm','-r','200','-png','-singlefile',str(out),str(out.with_suffix(''))],check=True)
    from pypdf import PdfReader
    pdf = PdfReader(out)
    if len(pdf.pages) != 1:
        raise ValueError('PDF exceeds one page: shorten before upload')
    text = pdf.pages[0].extract_text()
    if 'Yugal Teotia' not in text or 'AI-assisted' not in text or 'In progress' not in text:
        raise ValueError('PDF content validation failed')
    (folder / 'resume_text.txt').write_text(text)
    upsert({'Company':listing['company'],'Role':listing['role'],'Job ID':listing.get('job_id'),'Direct link':listing['url'],'Resume version':str(out.relative_to(ROOT)) + ' | SHA256 ' + hashlib.sha256(out.read_bytes()).hexdigest()}, 'truthful resume rendered')
    event('resume_rendered', slug=slug, pages=len(pdf.pages))

def prepare(slug, schema_path):
    check_pause()
    folder = job_folder(slug)
    if (folder/'application_stopped.json').exists():
        raise ValueError('Application stopped by user; explicit reauthorization required')
    listing = read_json(folder / 'listing.json')
    if listing and any(key(r)==key({'Company':listing.get('company'),'Job ID':listing.get('job_id'),'Role':listing.get('role'),'Direct link':listing.get('url')}) for r in read_json(STATE/'cleanup/other_history_keys.json',[])):
        raise ValueError('Historical application requisition already recorded; never reapply')
    if not listing or listing['fit'] != 'High':
        raise ValueError('Verified High fit required')
    fresh(listing['checked_at'])
    _,_,_, rows = tracker_rows()
    identity = {'Company':listing['company'],'Role':listing['role'],'Job ID':listing.get('job_id'),'Direct link':listing['url']}
    row = next(r for r in rows if key(r) == key(identity))
    if any(p in str(row.get('Status','')).lower() for p in PROTECTED):
        raise ValueError('Do not prepare an already actioned or uncertain application')
    verify_visual_review(folder)
    report = audit()
    schema = read_json(schema_path)
    if os.environ.get('JOB_SEARCH_WORKER_SLUG'):
        if not schema_path.resolve().is_relative_to(folder.resolve()):
            raise ValueError('Worker form schema must belong to its job folder')
        worker_url_guard(schema.get('url'))
    planned, blocked = [], []
    for field in schema.get('fields', []):
        answer = report['answers'].get(field['answer_key'], {'state':'uncertain'})
        if answer['state'] != 'confirmed':
            if field.get('required'):
                blocked.append(field['label'])
            continue
        if answer.get('mandatory_only') and not field.get('required'):
            continue
        planned.append({'label':field['label'],'answer_key':field['answer_key'],'value':answer['value'],'evidence':answer['evidence']})
    if schema.get('login_required'):
        blocked.append('Company ATS login handoff')
    result = {'time':now(),'url':schema['url'],'fields':planned,'blocked':blocked,'live_form_filled':False,
              'required_fields_verified':bool(schema.get('fields')) and not blocked,'submission_enabled':False,
              'stage':'Blocked' if blocked else 'Local draft prepared; browser filling still required'}
    write_json(folder / 'application_draft.json', result)
    upsert({**identity,'Status':'Blocked' if blocked else 'Draft saved','Next step':'; '.join(blocked) or ('Fill observed form, verify live values and complete pre-submit gate in this run' if read_json(STATE/'control.json',{}).get('submission_enabled') else 'Fill observed form, verify values, stop before Submit')}, 'application draft prepared')
    event('application_prepared', slug=slug, blocked=blocked, live_form_filled=False)
    return result

def match_mail(path):
    check_pause()
    import terminal_cleanup as tc
    messages = read_json(path)
    _,_,_, rows = tracker_rows()
    result=[]
    seen = {r.get('message_id') for r in read_json(STATE/'mail_matches.json',[])}
    for m in messages:
        if m['message_id'] in seen:
            continue
        # Only explicit requisition ID + company alias is eligible for automatic row linking.
        candidates=[r for r in rows if r.get('Job ID') and re.search(r'(?<!\w)'+re.escape(str(r['Job ID']))+r'(?!\w)', m['subject']) and employer(r['Company'])==employer(m.get('company',''))]
        item={'message_id':m['message_id'],'date':m['date'],'subject':m['subject'],'kind':m.get('kind','unknown'),'candidates':[key(r) for r in candidates], 'match':'exact' if len(candidates)==1 else 'review'}
        result.append(item)
        if len(candidates)==1:
            r=candidates[0]
            if tc.is_terminal(r.get('Status')):
                item['match']='terminal-existing'
                continue
            response=f"{m['date']} {m.get('kind','unknown')}: Gmail {m['message_id']}"
            previous=r.get('Responses') or ''
            if response not in previous:
                upsert({'Company':r['Company'],'Role':r['Role'],'Job ID':r['Job ID'],'Responses':(previous+'\n'+response).strip()}, 'Gmail response matched (no mail sent)')
            # An acknowledgment alone does not establish that this workflow submitted.
    write_json(STATE/'mail_matches.json',read_json(STATE/'mail_matches.json',[])+result)
    event('mail_matched', exact=sum(r['match']=='exact' for r in result), review=sum(r['match']=='review' for r in result))
    return result

def confirmation(slug, path):
    check_pause()
    evidence=read_json(path)
    if not evidence.get('application_id') or not evidence.get('captured_at') or not evidence.get('url'):
        raise ValueError('Confirmation requires application ID, capture time and source URL')
    worker_url_guard(evidence['url'])
    artifact=(ROOT/evidence['artifact']).resolve()
    if not artifact.is_relative_to(ROOT) or not artifact.is_file():
        raise ValueError('Confirmation artifact must exist inside workspace')
    folder=job_folder(slug)
    if os.environ.get('JOB_SEARCH_WORKER_SLUG') and not artifact.is_relative_to(folder.resolve()):
        raise ValueError('Worker confirmation proof must belong to its job folder')
    listing=read_json(folder/'listing.json')
    if str(evidence.get('job_id'))!=str(listing['job_id']):
        raise ValueError('Confirmation requisition mismatch')
    identity={'Company':listing['company'],'Job ID':listing['job_id']}
    if key(identity)==('unitedhealthgroup','2388999'):
        raise ValueError('Completed Optum requisition cannot be retried')
    write_json(folder/'confirmation.json',evidence)
    upsert({**identity,'Role':listing['role'],'Status':'Applied','Submission date':evidence['captured_at'],'Confirmation evidence':str((ROOT/'jobs'/slug/'confirmation.json').relative_to(ROOT))}, 'submission confirmation recorded',authorized_transition=True)

def begin_submission(slug,path):
    folder=job_folder(slug)
    listing=read_json(folder/'listing.json',{})
    if key({'Company':listing.get('company'),'Job ID':listing.get('job_id')})==('unitedhealthgroup','2388999'):
        raise ValueError('Never retry Optum 2388999')
    with (folder/'submission.lock').open('a') as lock:
        try: fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError: raise ValueError('Submission reservation already in progress; do not retry')
        return _begin_submission_locked(slug,path)

def _begin_submission_locked(slug,path):
    check_pause()
    if not read_json(STATE/'control.json',{}).get('submission_enabled'):
        raise ValueError('Submissions disabled')
    from daily_runner import gate
    gate(sys.modules[__name__])
    folder=job_folder(slug);listing=read_json(folder/'listing.json',{})
    identity={'Company':listing.get('company'),'Job ID':listing.get('job_id'),'Role':listing.get('role'),'Direct link':listing.get('url')}
    if any(key(r)==key(identity) for r in read_json(STATE/'cleanup/other_history_keys.json',[])):
        raise ValueError('Historical application requisition already recorded; never reapply')
    if key(identity)==('unitedhealthgroup','2388999'):
        raise ValueError('Never retry Optum 2388999')
    if not listing.get('job_id') or listing.get('fit')!='High' or not listing.get('original_verified') or not listing.get('live'):
        raise ValueError('Original live High-fit requisition required')
    if score(listing)[0]!='High':
        raise ValueError('High fit needs Java/Spring backend, realistic experience and India location evidence with no essential gap')
    fresh(listing['checked_at']);verify_visual_review(folder)
    row=next((r for r in tracker_rows()[3] if key(r)==key(identity)),{})
    if any(p in str(row.get('Status','')).lower() for p in PROTECTED) or (folder/'submission_attempt.json').exists():
        raise ValueError('Already actioned or uncertain application; investigate without resubmitting')
    review=read_json(path,{})
    required=['required_answers_confirmed','live_values_verified','resume_attachment_verified','history_deduplicated','certification_approved']
    if not all(review.get(k) is True for k in required):
        raise ValueError('Missing required pre-submit verification')
    for k in ['form_evidence','history_evidence','certification_evidence']:
        artifact=(ROOT/review.get(k,'')).resolve()
        if not artifact.is_relative_to(ROOT) or not artifact.is_file():
            raise ValueError('Missing local pre-submit evidence: '+k)
        if os.environ.get('JOB_SEARCH_WORKER_SLUG') and not artifact.is_relative_to(folder.resolve()):
            raise ValueError('Worker pre-submit evidence must belong to its job folder: '+k)
    fresh(review['verified_at'])
    # Persistent O_EXCL claim is shared by every slug/process for this requisition.
    # A crash leaves the claim in place for read-only investigation, never retry.
    claims=STATE/'submission-claims';claims.mkdir(parents=True,exist_ok=True)
    claim_name=hashlib.sha256(('\0'.join(key(identity))).encode()).hexdigest()+'.json'
    claim_path=claims/claim_name
    try:
        descriptor=os.open(claim_path,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
    except FileExistsError:
        raise ValueError('Requisition already has a persistent submission claim; do not retry')
    with os.fdopen(descriptor,'w') as claim_file:
        json.dump({'company_key':key(identity)[0],'job_id':listing['job_id'],'slug':slug,'url':listing.get('url'),'claimed_at':now()},claim_file)
        claim_file.flush();os.fsync(claim_file.fileno())
    # Write the uncertain state BEFORE a network action. Any interruption prevents retry.
    upsert({**identity,'Status':'Submission unconfirmed','Next step':'Submission authorized/in progress; verify result before any retry'},'pre-submit safety gate')
    write_json(folder/'submission_attempt.json',{'started_at':now(),'review':str(path.relative_to(ROOT)) if path.is_absolute() else str(path),'job_id':listing['job_id'],'state':'Pending confirmation; never retry automatically'})
    event('submission_attempt_registered',slug=slug)
    return {'authorized_to_click_once':True,'job_id':listing['job_id']}

def verify_visual_review(folder):
    review=read_json(folder/'resume_visual_review.json',{})
    if review.get('status')!='visually_clean':
        raise ValueError('Resume has no completed visual review')
    latest=read_json(folder/'resume_latest.json',{})
    if latest.get('pdf')!=review.get('pdf') or latest.get('pdf_sha256')!=review.get('pdf_sha256'):
        raise ValueError('Visual review does not match the latest resume')
    spec=folder/'resume_spec.json'
    if spec.exists() and hashlib.sha256(spec.read_bytes()).hexdigest()!=review.get('spec_sha256'):
        raise ValueError('Resume specification changed since visual review')
    for field in ['pdf','png']:
        path=(folder/review[field]).resolve()
        if not path.is_relative_to(folder.resolve()) or not path.is_file():
            raise ValueError('Visual review artifact missing or outside job folder')
        if hashlib.sha256(path.read_bytes()).hexdigest()!=review[field+'_sha256']:
            raise ValueError('Resume changed since visual review; render and inspect again')
    if not all(review.get('checks',{}).get(c) for c in ['letters','links','margins','bullets']):
        raise ValueError('Visual review incomplete')
    return review

def stop_application(slug):
    folder=job_folder(slug)
    listing=read_json(folder/'listing.json')
    write_json(folder/'application_stopped.json',{'stopped_at':now(),'reason':'User explicitly stopped Optum application work','submission_enabled':False})
    draft=read_json(folder/'application_draft.json',{})
    draft.update({'stage':'Stopped by user','live_form_filled':False,'submission_enabled':False})
    write_json(folder/'application_draft.json',draft)
    upsert({'Company':listing['company'],'Role':listing['role'],'Job ID':listing['job_id'],'Status':'Stopped','Resume version':'Previous PDF rejected; corrected v2 is for local review only','Next step':'Do not resume Optum application without explicit user authorization. Profile reconciliation and PDF repair only.'},'application stopped by user',allow_paused=True)
    event('application_stopped',slug=slug)

def browser_step(slug, path):
    check_pause()
    folder=job_folder(slug)
    if (folder/'application_stopped.json').exists():
        raise ValueError('Application stopped by user; no further browser steps authorized')
    step=read_json(path)
    allowed=['opened_application','login_blocked','fields_filled','resume_uploaded','values_verified','stopped_before_submit','captcha_blocked']
    if step.get('action') not in allowed or not step.get('url') or not step.get('captured_at'):
        raise ValueError('Browser step needs supported action, URL and capture timestamp')
    worker_url_guard(step['url'])
    artifact=(ROOT/step['artifact']).resolve()
    if not artifact.is_relative_to(ROOT) or not artifact.is_file():
        raise ValueError('Browser evidence must exist inside workspace')
    if os.environ.get('JOB_SEARCH_WORKER_SLUG') and not artifact.is_relative_to(folder.resolve()):
        raise ValueError('Worker browser evidence must belong to its job folder')
    safe={k:step[k] for k in ['action','url','captured_at','artifact']}
    history=read_json(folder/'browser_steps.json',[])
    if safe not in history:
        write_json(folder/'browser_steps.json',history+[safe])
    event('browser_step',slug=slug,**safe)
    if step['action']=='stopped_before_submit':
        draft=read_json(folder/'application_draft.json',{})
        verified=any(s['action']=='values_verified' for s in history)
        filled=any(s['action']=='fields_filled' for s in history)
        if verified and filled and draft.get('required_fields_verified'):
            verify_visual_review(folder)
            draft['live_form_filled']=True;draft['stage']='Ready to submit; disabled'
            write_json(folder/'application_draft.json',draft)
            listing=read_json(folder/'listing.json')
            upsert({'Company':listing['company'],'Role':listing['role'],'Job ID':listing['job_id'],'Status':'Ready to submit','Next step':'Submission disabled; reverify before any later action'},'browser draft verified; no submission')

def run(stage, execute=False):
    check_pause()
    if execute:
        import daily_runner as dr
        result=dr.run(sys.modules[__name__],stage=stage)
        print(json.dumps(result,indent=2))
        return
    prompt=(ROOT/'WORKFLOW_AGENT.md').read_text()+f'\nRun only stage: {stage}. Submission remains disabled.\n'
    STATE.mkdir(exist_ok=True)
    with (STATE/'run.lock').open('a') as lock:
        try: fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError: raise ValueError('A run is already active')
        stamp=dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S')
        request=STATE/'runs'/f'{stamp}-{stage}.md'
        request.parent.mkdir(exist_ok=True)
        request.write_text(prompt)
        event('run_requested',stage=stage,execute=execute)
        if not execute:
            print(f'Request saved: {request.relative_to(ROOT)}. Run with --execute to start local Codex CLI, or use this request in the current connected session.')
            return

def main():
    ap=argparse.ArgumentParser()
    sub=ap.add_subparsers(dest='cmd',required=True)
    for c in ['status','audit','start','pause','resume','logs','changes','cleanup-preview','cleanup-apply','cleanup-prune','redact-browser-artifacts','master-check','portals']:
        sub.add_parser(c)
    sub.add_parser('queue-status')
    p=sub.add_parser('fit');p.add_argument('job_id',type=int);p.add_argument('packet',type=Path)
    sub.add_parser('review')
    p=sub.add_parser('leads');p.add_argument('action',choices=['resolve','review-portal-native']);p.add_argument('--limit',type=int,default=None);p.add_argument('--target-ids',default=None,help='Comma-separated exact portal requisition identifiers for a targeted recheck')
    p=sub.add_parser('config');p.add_argument('action',choices=['check','naukri-live','linkedin-live','read-only']);p.add_argument('--scheduled',action='store_true')
    p=sub.add_parser('sources');p.add_argument('action',choices=['list'])
    p=sub.add_parser('db');p.add_argument('action',choices=['migrate','verify'])
    p=sub.add_parser('queue-cooldown');p.add_argument('slug');p.add_argument('reason',choices=['rate_limit','site_failure','verification_required']);p.add_argument('seconds',type=int)
    for c in ['ingest','match-mail','upsert']:
        sub.add_parser(c).add_argument('path',type=Path)
    for c in ['tailor','prepare','record-confirmation','record-browser-step','stop-application','begin-submission']:
        p=sub.add_parser(c);p.add_argument('slug')
        if c not in ('tailor','stop-application'):p.add_argument('path',type=Path)
    p=sub.add_parser('run');p.add_argument('stage',choices=['search','prepare','mail']);p.add_argument('--execute',action='store_true')
    p=sub.add_parser('daily');p.add_argument('--test',action='store_true');p.add_argument('--scheduled',action='store_true');p.add_argument('--max-seconds',type=int,default=1800);p.add_argument('--page-cap',type=int);p.add_argument('--query-cap',type=int);p.add_argument('--sources');p.add_argument('--review-limit',type=int);p.add_argument('--fresh-only',action='store_true')
    p=sub.add_parser('apply-now',help='Run now using the configured Naukri/LinkedIn execution modes');p.add_argument('--dry-run','--read-only',action='store_true');p.add_argument('--sources');p.add_argument('--fresh-only',action='store_true');p.add_argument('--review-limit',type=int);p.add_argument('--max-seconds',type=int);p.add_argument('--page-cap',type=int);p.add_argument('--query-cap',type=int);p.add_argument('--job-id',type=int,help='Restrict application work to this exact SQLite job');p.add_argument('--queue-only',action='store_true',help='Process saved candidates without discovery');p.add_argument('--session-check-only',action='store_true',help='Check the dedicated LinkedIn session only; no discovery or application actions')
    p=sub.add_parser('ui',help='Start the localhost dashboard');p.add_argument('--port',type=int,default=8765)
    p=sub.add_parser('pipeline');p.add_argument('--test',action='store_true');p.add_argument('--max-seconds',type=int,default=1800);p.add_argument('--page-cap',type=int);p.add_argument('--query-cap',type=int)
    p=sub.add_parser('needs-answer');p.add_argument('slug');p.add_argument('path',type=Path)
    sub.add_parser('reconcile').add_argument('path',type=Path)
    p=sub.add_parser('master-build',help='Build/check the generic master PDF from the .tex');p.add_argument('--install',action='store_true');p.add_argument('--visually-reviewed',action='store_true')
    p=sub.add_parser('alert-ack');p.add_argument('slug')
    p=sub.add_parser('alert-resolve');p.add_argument('slug');p.add_argument('path',type=Path)
    p=sub.add_parser('alerts');p.add_argument('action',choices=['test-desktop','phone-setup','test-phone','confirm-phone','status'])
    p=sub.add_parser('view');p.add_argument('--all',action='store_true')
    for c in ['backfill-preview','backfill-import']:
        sub.add_parser(c).add_argument('path',type=Path)
    args=ap.parse_args()
    try:
        import job_db
        if job_db.canonical_enabled() and args.cmd in ('ingest','upsert','tailor','prepare','match-mail','record-confirmation','begin-submission','record-browser-step','stop-application','needs-answer','alert-resolve','backfill-import','cleanup-apply','cleanup-prune','queue-cooldown'):
            raise ValueError('Legacy application mutation is disabled after SQLite migration; use guarded pipeline transitions')
        if job_db.canonical_enabled() and args.cmd=='run' and args.execute:
            raise ValueError('Legacy Codex execution is disabled after SQLite migration')
        if args.cmd=='config':
            import job_config
            if args.action=='linkedin-live':job_config.set_linkedin_mode(True,scheduled=args.scheduled)
            elif args.action!='check':job_config.set_naukri_mode(args.action=='naukri-live',scheduled=args.scheduled)
            config=job_config.load()
            print(json.dumps({'valid':True,'mode':config['mode'],'sources':len(job_config.source_rows(config)),
                'enabled_sources':sum(x['enabled'] for x in job_config.source_rows(config)),
                'submissions_enabled':config['submissions']['enabled'],
                'naukri_worker_enabled':config['naukri_worker']['enabled'],
                'manual_live_enabled':job_config.live_naukri_enabled(config),
                'timer_live_enabled':job_config.live_naukri_enabled(config,timer=True),
                'naukri_max_submissions_per_run':config['naukri_worker']['max_submissions_per_run'],
                'providers':{p:{'manual_configured':job_config.live_provider_enabled(config,p),'timer_configured':job_config.live_provider_enabled(config,p,timer=True),
                    'max_submissions_per_run':config['portal_workers'][p].get('max_submissions_per_run',1),
                    'max_new_attempts_per_run':config['portal_workers'][p].get('max_new_attempts_per_run',config['portal_workers'][p].get('max_submissions_per_run',1))} for p in ('naukri','linkedin')},
                'other_portals':'disabled'},indent=2))
        elif args.cmd=='sources':
            import job_config
            print(json.dumps(job_config.source_rows(job_config.load()),indent=2))
        elif args.cmd=='db':
            import job_db
            if args.action=='migrate':
                report=job_db.migrate()
                import job_views
                report['views']=job_views.regenerate()
                print(json.dumps(report,indent=2,default=str))
            else:
                report=read_json(STATE/'migration/verification.json',{})
                db=job_db.connect()
                report['current_rows']=db.execute('SELECT count(*) FROM jobs').fetchone()[0]
                report['current_history_keys']=db.execute('SELECT count(*) FROM history_keys').fetchone()[0]
                report['current_integrity']=db.execute('PRAGMA integrity_check').fetchone()[0]
                db.close();print(json.dumps(report,indent=2,default=str))
        elif args.cmd in ['view','backfill-preview','backfill-import']:
            import tracker_views as tv
            if args.cmd=='view' and (STATE/'canonical-db-enabled.json').is_file():
                import job_views
                result=job_views.regenerate()
            else:
                result = tv.views(sys.modules[__name__],args.all) if args.cmd=='view' else tv.preview(sys.modules[__name__],args.path) if args.cmd=='backfill-preview' else tv.import_reviewed(sys.modules[__name__],args.path)
            print(json.dumps(result,indent=2,default=str))
        elif args.cmd=='audit':print(json.dumps(audit(),indent=2))
        elif args.cmd=='ui':
            import ui_server
            ui_server.serve(args.port)
        elif args.cmd=='portals':
            import job_config,registry_cli
            print(json.dumps(registry_cli.status(job_config.load()),indent=2))
        elif args.cmd=='master-check':
            import master_check
            print(json.dumps(master_check.check(),indent=2))
        elif args.cmd=='master-build':
            import master_build
            print(json.dumps(master_build.build(install=args.install,visually_reviewed=args.visually_reviewed),indent=2))
        elif args.cmd=='reconcile':
            import response_reconcile
            print(json.dumps(response_reconcile.run(args.path),indent=2))
        elif args.cmd=='needs-answer':
            import urgent_alerts as alerts
            print(json.dumps(alerts.mark(sys.modules[__name__],args.slug,args.path),indent=2))
        elif args.cmd=='alert-ack':
            import urgent_alerts as alerts
            print(json.dumps(alerts.acknowledge(sys.modules[__name__],args.slug),indent=2))
        elif args.cmd=='alert-resolve':
            import urgent_alerts as alerts
            print(json.dumps(alerts.resolve(sys.modules[__name__],args.slug,args.path),indent=2))
        elif args.cmd=='alerts':
            import urgent_alerts as alerts
            result=alerts.test_desktop() if args.action=='test-desktop' else alerts.setup_phone() if args.action=='phone-setup' else alerts.test_phone(sys.modules[__name__]) if args.action=='test-phone' else alerts.confirm_phone(sys.modules[__name__]) if args.action=='confirm-phone' else alerts.status(sys.modules[__name__])
            print(json.dumps(result,indent=2))
        elif args.cmd.startswith('cleanup-'):
            import terminal_cleanup as tc
            result=tc.preview(sys.modules[__name__]) if args.cmd=='cleanup-preview' else tc.apply(sys.modules[__name__]) if args.cmd=='cleanup-apply' else tc.prune_retained(sys.modules[__name__])
            print(json.dumps(result,indent=2))
        elif args.cmd=='redact-browser-artifacts':
            count=0
            for path in (ROOT/'.playwright-mcp').glob('*.yml'):
                text=path.read_text()
                sanitized=re.sub(r'(?mi)^(.*?textbox[^\n]*?"[^\n]*?password[^\n]*?"):\s*[^\n]*$',r'\1: [REDACTED]',text)
                if sanitized!=text:
                    path.write_text(sanitized);count+=1
            event('browser_artifacts_redacted',files=count);print(f'Redacted password fields in {count} browser snapshots')
        elif args.cmd=='start':
            import job_config
            cfg=job_config.load()
            if not cfg['schedule']['enabled']:raise ValueError('Schedule disabled in config')
            import control
            control.resume(cfg)
            print('Timer active: '+('Naukri live, max one application per run.' if job_config.live_naukri_enabled(cfg,timer=True) else 'read-only discovery.'))
        elif args.cmd=='daily':
            import pipeline_v2
            print(json.dumps(pipeline_v2.run(test=args.test,scheduled=args.scheduled,max_seconds=args.max_seconds,
                                             page_cap=args.page_cap,query_cap=args.query_cap,
                                             sources=tuple(x.strip() for x in args.sources.split(',') if x.strip()) if args.sources else None,
                                             review_limit=args.review_limit,fresh_only=args.fresh_only),indent=2))
        elif args.cmd=='apply-now':
            import apply_now
            outcome=apply_now.run(dry_run=args.dry_run,sources=tuple(args.sources.split(',')) if args.sources else None,
                fresh_only=args.fresh_only,review_limit=args.review_limit,max_seconds=args.max_seconds,page_cap=args.page_cap,query_cap=args.query_cap,application_job_id=args.job_id,queue_only=args.queue_only,session_check_only=args.session_check_only)
            text,started=apply_now.summarize(outcome)
            print('\n'+text)
            if not started:sys.exit(2)
        elif args.cmd=='pipeline':
            import pipeline_v2
            print(json.dumps(pipeline_v2.run(test=args.test,max_seconds=args.max_seconds,page_cap=args.page_cap,query_cap=args.query_cap),indent=2))
        elif args.cmd=='queue-status':
            if job_db.canonical_enabled():
                db=job_db.connect()
                print(json.dumps([dict(r) for r in db.execute('SELECT q.*,j.employer,j.title,j.requisition_id FROM queue q JOIN jobs j ON j.id=q.job_id ORDER BY q.priority DESC,q.queued_at')],indent=2))
                db.close()
            else:
                import priority_queue
                print(json.dumps(priority_queue.status(),indent=2))
        elif args.cmd=='fit':
            import job_config
            if not job_db.canonical_enabled():raise ValueError('Canonical SQLite migration required')
            packet=json.loads(args.packet.read_text())
            db=job_db.connect()
            try:print(json.dumps(job_db.apply_fit_decision(db,args.job_id,packet,job_config.load()),indent=2))
            finally:db.close()
        elif args.cmd=='review':
            import job_config,review_candidates,job_views
            cfg=job_config.load();db=job_db.connect();job_db.initialize(db)
            try:
                ids=[r['id'] for r in db.execute("SELECT id FROM jobs WHERE state='Verified' AND jd_text IS NOT NULL")]
                added=review_candidates.enqueue(db,ids)
                checked=review_candidates.consume(db,cfg,cfg['review']['max_jds_per_run'],cfg['review']['per_company_cap'])
                by_source={};by_company={}
                for item in checked:
                    source=db.execute('SELECT source FROM jobs WHERE id=?',(item['job_id'],)).fetchone()['source']
                    by_source[source]=by_source.get(source,0)+1
                    by_company[item['company']]=by_company.get(item['company'],0)+1
                excluded=[];excluded_by_company={}
                candidates=db.execute("SELECT * FROM jobs WHERE state='Verified' AND jd_text IS NOT NULL AND availability='Live'").fetchall()
                for row in sorted(candidates,key=review_candidates.priority,reverse=True):
                    if len(excluded)>=5:break
                    if review_candidates.eligible(row) or not review_candidates.INDIA.search(row['location'] or ''):continue
                    if not re.search(r'engineer|developer',row['title'] or '',re.I):continue
                    company=row['employer']
                    if excluded_by_company.get(company,0)>=3:continue
                    excluded.append({'job_id':row['id'],'company':company,'role':row['title'],'source':row['source'],
                                     'direct_url':row['canonical_url'],'stage':'Hard filtered',
                                     'reasons':review_candidates.hard_filter_reasons(row)})
                    excluded_by_company[company]=excluded_by_company.get(company,0)+1
                selection=review_candidates.coverage(db)
                pending=db.execute("SELECT count(*) FROM review_queue WHERE state='Pending'").fetchone()[0]
                report={'mode':'read_only','queued_for_review':added,'reviewed':len(checked),'by_source':by_source,
                        'by_company':by_company,'high':sum(x['fit']=='High' for x in checked),'queued_for_resume':db.execute("SELECT count(*) FROM queue WHERE state='QueuedResume'").fetchone()[0],
                        'selection_coverage':selection,'pending_after_run':pending,
                        'roles':checked,'top_hard_filtered':excluded}
            finally:db.close()
            path=STATE/'reviews'/(dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S')+'-fair.json')
            write_json(path,report);report['log']=str(path.relative_to(ROOT));report['views']=job_views.regenerate()
            print(json.dumps(report,indent=2))
        elif args.cmd=='leads':
            import job_config,lead_resolution
            if args.action=='review-portal-native' and not args.target_ids:
                raise ValueError('review-portal-native requires exact --target-ids; broad portal searches are not run')
            cfg=job_config.load();db=job_db.connect();job_db.initialize(db)
            try:
                ids=[x for x in args.target_ids.split(',') if x] if args.target_ids else None
                if args.action=='review-portal-native':
                    report=lead_resolution.review_portal_native(db,cfg,ids,args.limit if args.limit is not None else 29)
                else:
                    report=lead_resolution.resolve(db,cfg,args.limit if args.limit is not None else cfg['review']['max_portal_leads_per_run'],ids)
            finally:db.close()
            path=STATE/'reviews'/(dt.datetime.now(dt.timezone.utc).strftime('%Y%m%dT%H%M%S')+('-portal-native' if args.action=='review-portal-native' else '-leads')+'.json')
            write_json(path,report);report['log']=str(path.relative_to(ROOT))
            if args.action=='review-portal-native':
                import job_views
                report['views']=job_views.regenerate()
            print(json.dumps(report,indent=2))
        elif args.cmd=='queue-cooldown':
            import priority_queue
            listing=read_json(job_folder(args.slug)/'listing.json',{})
            hostname=priority_queue.host(listing['url'])
            priority_queue.cooldown(hostname,args.seconds,args.reason)
            event('ats_cooldown',slug=args.slug,hostname=hostname,reason=args.reason,seconds=args.seconds)
            print(json.dumps({'host':hostname,'reason':args.reason,'seconds':args.seconds}))
        elif args.cmd in ('pause','resume'):
            import control,job_config
            if args.cmd=='pause':
                control.pause();print('Paused; timer stopped and running work stops at its next gate.')
            else:
                outcome=control.resume(job_config.load())
                print('Resumed configured timer.' if outcome['timer'] else 'Resumed manual work; timer remains off.')
            event(args.cmd)
        elif args.cmd=='status':
            import daily_runner as dr
            import urgent_alerts as alerts
            import priority_queue
            import job_db,job_config
            rows=job_db.tracker_rows() if job_db.canonical_enabled() else tracker_rows()[3]
            if job_db.canonical_enabled():
                db=job_db.connect();queue_counts={r['state']:r['n'] for r in db.execute('SELECT state,count(*) n FROM queue GROUP BY state')};db.close()
            else:queue_counts=priority_queue.status()['counts']
            cfg=job_config.load()
            print(json.dumps({'control':read_json(STATE/'control.json',{'paused':False,'submission_enabled':False,'schedule_enabled':False}),'canonical_database':job_db.canonical_enabled(),'config_mode':cfg['mode'],'concurrency':cfg['concurrency'],'queue_counts':queue_counts,'schedule':dr.schedule_status(),'alerts':alerts.status(sys.modules[__name__]),'last_daily_run':read_json(STATE/'last-daily-run.json',{}),'active':read_json(STATE/'active.json',{}),'rows':[{k:r.get(k) for k in ['Company','Role','Job ID','Fit','Status','Next step']} for r in rows]},indent=2))
        elif args.cmd in ('logs','changes'):
            if args.cmd=='logs':
                last=read_json(STATE/'active.json',{}).get('log') or read_json(STATE/'last-daily-run.json',{}).get('log')
                print((ROOT/last).read_text() if last else (STATE/'events.jsonl').read_text() if (STATE/'events.jsonl').exists() else 'No events')
            else:
                for p in sorted((STATE/'changes').glob('*.json')):print(p.name+'\n'+p.read_text())
        elif args.cmd=='ingest':print(ingest(args.path))
        elif args.cmd=='upsert':print(upsert(read_json(args.path),'reviewed tracker correction'))
        elif args.cmd=='tailor':tailor(args.slug)
        elif args.cmd=='prepare':print(json.dumps(prepare(args.slug,args.path),indent=2))
        elif args.cmd=='match-mail':print(json.dumps(match_mail(args.path),indent=2))
        elif args.cmd=='record-confirmation':confirmation(args.slug,args.path)
        elif args.cmd=='begin-submission':print(json.dumps(begin_submission(args.slug,args.path),indent=2))
        elif args.cmd=='record-browser-step':browser_step(args.slug,args.path)
        elif args.cmd=='stop-application':stop_application(args.slug)
        elif args.cmd=='run':run(args.stage,args.execute)
    except (ValueError,KeyError,FileNotFoundError) as exc:
        event('blocked',command=args.cmd,reason=str(exc));ap.exit(1,str(exc)+'\n')

if __name__=='__main__':main()
